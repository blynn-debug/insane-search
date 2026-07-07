"""
탱크옥션 건물등기 스크래핑 - 개인채권자 필터링
전체 데이터 페이지네이션 조회
"""

import requests
from bs4 import BeautifulSoup
import time
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

# ============================================
# 쿠키 설정 (브라우저에서 복사)
# ============================================
COOKIE_STRING = '_fwb=254J3uteonjfGwGBar1Z26W.1762748083804; _gcl_au=1.1.1315489470.1762748084; _wp_uid=1-REDACTED-s1762748079.221471|windows_10|chrome-cgzyhv; _ga=GA1.1.1953926687.1762748084; PHPSESSID=REDACTED_SESSION; CK_PD_VHIT=a%3A3%3A%7Bi%3A0%3Bs%3A7%3A%222383411%22%3Bi%3A1%3Bs%3A7%3A%222360378%22%3Bi%3A2%3Bs%3A7%3A%222319904%22%3B%7D; TKC_ABUSER=0; CK_PD_HIT=a%3A3%3A%7Bi%3A0%3Bs%3A7%3A%222331009%22%3Bi%3A1%3Bs%3A7%3A%222459779%22%3Bi%3A2%3Bs%3A7%3A%222336798%22%3B%7D; wcs_bt=3da52d33d5fe98:1770004848; _ga_4K6G99FPPX=GS2.1.s1770004819$o96$g1$t1770004849$j30$l0$h0'

# 금융기관 키워드 (제외)
FINANCIAL_KEYWORDS = ['은행', '카드', '캐피탈', '금고', '대부']

# 페이지당 조회 개수
PAGE_SIZE = 100

session = requests.Session()


def get_auction_list_page(page_no=1):
    """경매 목록 1페이지 조회 (POST)"""
    url = f'https://tankauction.com/ca/AuctList.php?srchCase=srchAll&pageNo={page_no}&dataSize={PAGE_SIZE}&pageSize=10'

    headers = {
        'Accept': 'application/json, text/javascript, */*; q=0.01',
        'Accept-Encoding': 'gzip, deflate, br, zstd',
        'Accept-Language': 'ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7,ja-JP;q=0.6,ja;q=0.5',
        'Connection': 'keep-alive',
        'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
        'Cookie': COOKIE_STRING,
        'Host': 'tankauction.com',
        'Origin': 'https://tankauction.com',
        'Referer': 'https://tankauction.com/ca/caList.php?page=1',
        'Sec-Fetch-Dest': 'empty',
        'Sec-Fetch-Mode': 'cors',
        'Sec-Fetch-Site': 'same-origin',
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Safari/537.36',
        'X-Requested-With': 'XMLHttpRequest',
        'sec-ch-ua': '"Not:A-Brand";v="99", "Google Chrome";v="145", "Chromium";v="145"',
        'sec-ch-ua-mobile': '?0',
        'sec-ch-ua-platform': '"Windows"',
    }

    # Form Data - dataSize도 동적으로 변경
    form_data = f'siCd=11&guCd=0&dnCd=0&dptCd=0&addr_cs_key=0&adrPlural=&adrPlural_cnt=0&adrsEtcSelect=0&adrsEtc=&sn1=0&sn2=&pn=&chkGrpCtgr=10&chkEaCtgr=201013&chkEaCtgr=201014&chkEaCtgr=201015%2C201017%2C201021&chkEaCtgr=201020&chkEaCtgr=201010&chkEaCtgr=201011%2C201012&chkEaCtgr=201022&chkEaCtgr=201016&chkEaCtgr=201018&stat=11&fbCntBgn=0&fbCntEnd=0&bgnDt=&endDt=&apslAmtBgn=0&apslAmtEnd=0&landSqmBgn=&landSqmEnd=&minbAmtBgn=0&minbAmtEnd=0&bldgSqmBgn=&bldgSqmEnd=&totFlrBgn=0&totFlrEnd=0&prsvBgn=0&prsvEnd=0&flrBgn=0&flrEnd=0&preBgnDt=&preEndDt=&dpslDvsn=0&auctType=0&minbPctBgn=0&minbPctEnd=0&maxPnBgn=0&maxPnEnd=0&local=0&line=0&station=0&distance=0&splSrchType=0&powerCtgrs=1&chkCtgrsCd=201013%7C201014%7C201015%2C201017%2C201021%7C201020%7C201010%7C201011%2C201012%7C201022%7C201016%7C201018&chkSplCdtn=&chkPrpsCdtn=&dataSize={PAGE_SIZE}&lsType=0&odrCol=14&odrAds=0&srchFR=0&idxFR=0&ck_photo=0'

    try:
        response = session.post(url, headers=headers, data=form_data)
        if response.status_code == 200 and response.text:
            return response.json()
    except Exception as e:
        print(f"API 호출 실패 (페이지 {page_no}): {e}")
    return None


def get_all_auction_list():
    """전체 경매 목록 페이지네이션으로 조회"""
    all_items = []
    page = 1

    while True:
        print(f"  페이지 {page} 조회 중...", end=" ")
        data = get_auction_list_page(page)

        if not data or 'item' not in data:
            print("응답 없음 - 종료")
            break

        items = data['item']
        if len(items) == 0:
            print("데이터 없음 - 종료")
            break

        all_items.extend(items)
        print(f"{len(items)}개 로드 (누적: {len(all_items)}개)")

        # 마지막 페이지 체크 (받아온 개수가 PAGE_SIZE보다 적으면 마지막)
        if len(items) < PAGE_SIZE:
            print("  마지막 페이지 도달")
            break

        page += 1
        time.sleep(0.3)

    return all_items


def get_registry_info(tid, total_count=100):
    """2차: 상세 페이지에서 건물등기 스크래핑 (GET)"""
    url = f'https://tankauction.com/ca/caView.php?tid={tid}&chkNo=1&TotNo={total_count}'

    headers = {
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7',
        'Accept-Encoding': 'gzip, deflate, br, zstd',
        'Accept-Language': 'ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7,ja-JP;q=0.6,ja;q=0.5',
        'Cache-Control': 'max-age=0',
        'Connection': 'keep-alive',
        'Cookie': COOKIE_STRING,
        'Host': 'tankauction.com',
        'Referer': 'https://tankauction.com/ca/caList.php?page=1',
        'Sec-Fetch-Dest': 'document',
        'Sec-Fetch-Mode': 'navigate',
        'Sec-Fetch-Site': 'same-origin',
        'Sec-Fetch-User': '?1',
        'Upgrade-Insecure-Requests': '1',
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Safari/537.36',
        'sec-ch-ua': '"Not:A-Brand";v="99", "Google Chrome";v="145", "Chromium";v="145"',
        'sec-ch-ua-mobile': '?0',
        'sec-ch-ua-platform': '"Windows"',
    }

    try:
        response = session.get(url, headers=headers, timeout=30)
        if response.status_code != 200:
            return None
    except Exception as e:
        print(f"  요청 오류: {e}")
        return None

    soup = BeautifulSoup(response.text, 'html.parser')

    registry_section = soup.find('div', {'id': 'lyCnt_regist'})
    if not registry_section:
        return {'채권합계': 0, '등기목록': [], '개인채권자있음': False}

    result = {'채권합계': 0, '등기목록': [], '개인채권자있음': False}

    # 채권합계금액
    span_box = registry_section.find('span', class_='spanBox')
    if span_box:
        amount_span = span_box.find('span', class_='selectedNumType')
        if amount_span and amount_span.get('data-value'):
            try:
                result['채권합계'] = int(amount_span.get('data-value'))
            except:
                result['채권합계'] = 0

    # 건물등기 테이블
    table = registry_section.find('table', class_='Ltbl_list')
    if not table:
        return result

    rows = table.find_all('tr')
    for row in rows:
        tds = row.find_all('td')
        if len(tds) >= 7:
            order = tds[0].get_text(strip=True)

            date_text = tds[1].get_text(strip=True)
            date_match = re.search(r'(\d{4}-\d{2}-\d{2})', date_text)
            receipt_date = date_match.group(1) if date_match else date_text

            receipt_no_span = tds[1].find('span', class_='rcNo')
            receipt_no = ''
            if receipt_no_span:
                receipt_no = receipt_no_span.get_text(strip=True).replace('(', '').replace(')', '')

            right_type = tds[2].get_text(strip=True)
            creditor = tds[3].get_text(strip=True)

            amount = 0
            amount_span = tds[4].find('span', class_='selectedNumType')
            if amount_span and amount_span.get('data-value'):
                try:
                    amount = int(amount_span.get('data-value'))
                except:
                    amount = 0

            note = tds[5].get_text(separator=' ', strip=True)
            extinction = tds[6].get_text(strip=True)

            # 개인채권자 여부 (금융기관 키워드 미포함)
            is_individual = True
            for keyword in FINANCIAL_KEYWORDS:
                if keyword in creditor:
                    is_individual = False
                    break

            entry = {
                '순서': order,
                '접수일': receipt_date,
                '접수번호': receipt_no,
                '권리종류': right_type,
                '권리자': creditor,
                '채권금액': amount,
                '비고': note,
                '소멸': extinction,
                '개인채권자': '예' if is_individual else '아니오'
            }
            result['등기목록'].append(entry)

            if is_individual and right_type in ['근저당권설정', '전세권설정', '주택임차권', '임차권등기']:
                result['개인채권자있음'] = True

    return result


def create_excel(all_results, output_path):
    """엑셀 파일 생성"""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = '개인채권자물건'
    ws2 = wb.create_sheet('등기상세')

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    individual_fill = PatternFill('solid', fgColor='FFFF00')
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    headers1 = ['TID', '사건번호', '물건종류', '소재지', '감정가', '최저가', '진행상태', '입찰일', '채권합계', '개인채권자목록']
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    headers2 = ['TID', '사건번호', '순서', '접수일', '접수번호', '권리종류', '권리자', '채권금액', '비고', '소멸', '개인채권자']
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    row1 = 2
    row2 = 2

    for item in all_results:
        if item['개인채권자있음']:
            individual_list = [e['권리자'] for e in item['등기목록']
                             if e['개인채권자'] == '예' and e['권리종류'] in ['근저당권설정', '전세권설정', '주택임차권', '임차권등기']]

            ws1.cell(row=row1, column=1, value=item['tid']).border = border
            ws1.cell(row=row1, column=2, value=item['사건번호']).border = border
            ws1.cell(row=row1, column=3, value=item['물건종류']).border = border
            ws1.cell(row=row1, column=4, value=item['소재지']).border = border
            ws1.cell(row=row1, column=5, value=item['감정가']).border = border
            ws1.cell(row=row1, column=6, value=item['최저가']).border = border
            ws1.cell(row=row1, column=7, value=item['진행상태']).border = border
            ws1.cell(row=row1, column=8, value=item['입찰일']).border = border
            ws1.cell(row=row1, column=9, value=item['채권합계']).border = border
            ws1.cell(row=row1, column=10, value=', '.join(individual_list)).border = border
            row1 += 1

        for entry in item['등기목록']:
            ws2.cell(row=row2, column=1, value=item['tid']).border = border
            ws2.cell(row=row2, column=2, value=item['사건번호']).border = border
            ws2.cell(row=row2, column=3, value=entry['순서']).border = border
            ws2.cell(row=row2, column=4, value=entry['접수일']).border = border
            ws2.cell(row=row2, column=5, value=entry['접수번호']).border = border
            ws2.cell(row=row2, column=6, value=entry['권리종류']).border = border

            cell_creditor = ws2.cell(row=row2, column=7, value=entry['권리자'])
            cell_creditor.border = border
            if entry['개인채권자'] == '예':
                cell_creditor.fill = individual_fill

            ws2.cell(row=row2, column=8, value=entry['채권금액']).border = border
            ws2.cell(row=row2, column=9, value=entry['비고']).border = border
            ws2.cell(row=row2, column=10, value=entry['소멸']).border = border
            ws2.cell(row=row2, column=11, value=entry['개인채권자']).border = border
            row2 += 1

    for col, w in zip('ABCDEFGHIJ', [10,15,12,50,15,15,10,12,15,30]):
        ws1.column_dimensions[col].width = w
    for col, w in zip('ABCDEFGHIJK', [10,15,8,12,10,15,20,15,30,8,12]):
        ws2.column_dimensions[col].width = w

    ws1.freeze_panes = 'A2'
    ws2.freeze_panes = 'A2'
    wb.save(output_path)


def main():
    print("=" * 60)
    print("탱크옥션 건물등기 스크래핑 - 개인채권자 필터링")
    print("=" * 60)

    # 1차: 전체 경매 목록 페이지네이션으로 조회
    print("\n[1단계] 경매 목록 전체 조회")
    items = get_all_auction_list()

    if not items:
        print("경매 목록 조회 실패! 쿠키를 확인하세요.")
        return

    total_count = len(items)
    print(f"\n총 {total_count}개 물건 발견\n")

    # 2차: 각 물건의 건물등기 스크래핑
    print("[2단계] 건물등기 상세 스크래핑")
    all_results = []
    individual_count = 0

    for idx, item in enumerate(items, 1):
        tid = item['tid']
        sa_no = item.get('saNo', '')
        addr = item.get('regnAdrs', '')
        category = item.get('ctgr', '')
        apsl_amt = item.get('apslAmt', 0)
        minb_amt = item.get('minbAmt', 0)
        stat_nm = item.get('statNm', '')
        bid_dt = item.get('bidDt', '')

        print(f"[{idx}/{total_count}] {sa_no} (tid={tid})", end=" ")

        registry_info = get_registry_info(tid, total_count)

        if registry_info:
            result = {
                'tid': tid, '사건번호': sa_no, '물건종류': category, '소재지': addr,
                '감정가': apsl_amt, '최저가': minb_amt, '진행상태': stat_nm, '입찰일': bid_dt,
                '채권합계': registry_info.get('채권합계', 0),
                '등기목록': registry_info.get('등기목록', []),
                '개인채권자있음': registry_info.get('개인채권자있음', False)
            }
            all_results.append(result)

            if result['개인채권자있음']:
                individual_count += 1
                print(f"-> 등기 {len(registry_info['등기목록'])}건 [★개인채권자]")
            else:
                print(f"-> 등기 {len(registry_info['등기목록'])}건")
        else:
            print("-> 정보없음")

        time.sleep(0.3)

    # 엑셀 저장
    output_path = '건물등기_개인채권자.xlsx'
    create_excel(all_results, output_path)

    print("\n" + "=" * 60)
    print(f"완료!")
    print(f"- 전체 물건: {len(all_results)}개")
    print(f"- 개인채권자 포함: {individual_count}개")
    print(f"- 파일: {output_path}")
    print("=" * 60)


if __name__ == '__main__':
    main()