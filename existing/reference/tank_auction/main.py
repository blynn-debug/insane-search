"""
탱크옥션 임차인 현황 스크래핑
사용법: python tankauction_scraper.py
"""

import requests
from bs4 import BeautifulSoup
import pandas as pd
import time
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============================================
# 쿠키 설정 (브라우저에서 복사한 문자열 그대로 붙여넣기)
# ============================================
COOKIE_STRING = '_fwb=254J3uteonjfGwGBar1Z26W.1762748083804; _gcl_au=1.1.1315489470.1762748084; _wp_uid=1-REDACTED-s1762748079.221471|windows_10|chrome-cgzyhv; _ga=GA1.1.1953926687.1762748084; PHPSESSID=REDACTED_SESSION; CK_PD_VHIT=a%3A3%3A%7Bi%3A0%3Bs%3A7%3A%222383411%22%3Bi%3A1%3Bs%3A7%3A%222360378%22%3Bi%3A2%3Bs%3A7%3A%222319904%22%3B%7D; TKC_ABUSER=0; wcs_bt=3da52d33d5fe98:1769998027; CK_PD_HIT=a%3A3%3A%7Bi%3A0%3Bs%3A7%3A%222146500%22%3Bi%3A1%3Bs%3A7%3A%222331009%22%3Bi%3A2%3Bs%3A7%3A%222459779%22%3B%7D; _ga_4K6G99FPPX=GS2.1.s1769997981$o95$g1$t1769998139$j60$l0$h0'

# 첨부된 JSON 데이터 (1차 API 응답 대체)
AUCTION_DATA = {"item":[{"tid":2146500,"ctgr":"다세대주택","saNo":"2022-57388","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 천호동 221-62, 비동 3층303호 (천호동,씨팰리스4차)","apslAmt":463000000,"minbAmt":16290000,"statNm":"유찰 15회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인"},{"tid":2331009,"ctgr":"도시형생활주택","saNo":"2024-2566(1)","crtDpt":"서울동부4계","regnAdrs":"서울 광진구 중곡동 56-18, 3층 301호 외 1필지","apslAmt":381000000,"minbAmt":156058000,"statNm":"유찰 3회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2459779,"ctgr":"다가구주택","saNo":"2024-56433","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 군자동 159-7","apslAmt":2097912920,"minbAmt":1342664000,"statNm":"유찰 2회","bidDt":"26.02.02","splCdtn":"임차권등기"},{"tid":2302692,"ctgr":"오피스텔(주거)","saNo":"2024-56945","crtDpt":"서울동부4계","regnAdrs":"서울 송파구 송파동 85, 6층603호 (송파동,에스아이팰리스레이크잠실)","apslAmt":287000000,"minbAmt":229600000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2323076,"ctgr":"다세대주택","saNo":"2024-58491","crtDpt":"서울동부4계","regnAdrs":"서울 송파구 가락동 151-4, 4층402호","apslAmt":260000000,"minbAmt":85197000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"위반건축물,임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2323077,"ctgr":"도시형생활주택","saNo":"2024-58606","crtDpt":"서울동부4계","regnAdrs":"서울 송파구 문정동 80-8, 4층401호 (문정동,집앤사베스타) 외 1필지","apslAmt":338000000,"minbAmt":270400000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2317247,"ctgr":"오피스텔(주거)","saNo":"2024-58736","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 둔촌동 435-1, 7층702호 (둔촌동,한울둔촌베아체) 외 1필지","apslAmt":338000000,"minbAmt":110756000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 3~4억"},{"tid":2317250,"ctgr":"오피스텔(주거)","saNo":"2024-59005","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 둔촌동 435-1, 5층504호 (둔촌동,한울둔촌베아체) 외 1필지","apslAmt":323000000,"minbAmt":258400000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 3~4억,HUG 임차권 인수조건변경"},{"tid":2328912,"ctgr":"도시형생활주택","saNo":"2024-60470","crtDpt":"서울동부4계","regnAdrs":"서울 송파구 방이동 151, 2층 201호 (방이동,호수(가)하우스)","apslAmt":390000000,"minbAmt":312000000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 2~3억,HUG 임차권 인수조건변경"},{"tid":2334768,"ctgr":"상가주택","saNo":"2024-61206","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 천호동 538-7","apslAmt":3141712740,"minbAmt":2513370000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 2~3억"},{"tid":2336796,"ctgr":"오피스텔(주거)","saNo":"2024-61220","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 둔촌동 435-1, 4층 402호 (둔촌동,한울둔촌베아체) 외 1필지","apslAmt":336000000,"minbAmt":268800000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 3~4억,HUG 임차권 인수조건변경"},{"tid":2336798,"ctgr":"도시형생활주택","saNo":"2024-61367","crtDpt":"서울동부4계","regnAdrs":"서울 광진구 구의동 251-177, 10층 1005호 (구의동,솔렌시아) 외 1필지","apslAmt":135000000,"minbAmt":135000000,"statNm":"유찰 4회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 1억이하,HUG 임차권 인수조건변경"},{"tid":2338575,"ctgr":"오피스텔(주거)","saNo":"2024-61626","crtDpt":"서울동부4계","regnAdrs":"서울 성동구 홍익동 138, 2층 201호 (홍익동,뉴센트빌) 외 2필지","apslAmt":331000000,"minbAmt":211840000,"statNm":"유찰 4회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 3~4억,HUG 임차권 인수조건변경"},{"tid":2338576,"ctgr":"다세대주택","saNo":"2024-61701","crtDpt":"서울동부4계","regnAdrs":"서울 광진구 구의동 23-31, 5층 501호","apslAmt":236500000,"minbAmt":189200000,"statNm":"유찰 5회","bidDt":"26.02.02","splCdtn":"위반건축물,임차권등기,공시가 1~2억,HUG 임차권 인수조건변경"},{"tid":2348579,"ctgr":"도시형생활주택","saNo":"2024-62483","crtDpt":"서울동부4계","regnAdrs":"서울 광진구 중곡동 191-13, 3층 301호 (중곡동,파크홈2차) 외 2필지","apslAmt":326000000,"minbAmt":326000000,"statNm":"유찰 4회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 2~3억,HUG 임차권 인수조건변경"},{"tid":2348585,"ctgr":"도시형생활주택","saNo":"2024-63080","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 성내동 245-20, 7층 702호 (성내동,화영하우스7차)","apslAmt":223000000,"minbAmt":142720000,"statNm":"유찰 2회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2348586,"ctgr":"도시형생활주택","saNo":"2024-63196","crtDpt":"서울동부4계","regnAdrs":"서울 송파구 가락동 12, 비동 3층 303호 (가락동,이든하우스)","apslAmt":300000000,"minbAmt":300000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2350725,"ctgr":"다세대주택","saNo":"2024-63301","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 성내동 46-22, 2층 202호 (성내동,더존예가비)","apslAmt":374000000,"minbAmt":153190000,"statNm":"유찰 4회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 3~4억"},{"tid":2353966,"ctgr":"다세대주택","saNo":"2024-63851","crtDpt":"서울동부4계","regnAdrs":"서울 성동구 사근동 230-1, 101동 4층 401호 (사근동,청계더하임)","apslAmt":231000000,"minbAmt":231000000,"statNm":"유찰 4회","bidDt":"26.02.02","splCdtn":"임차권등기,공시가 1~2억,HUG 임차권 인수조건변경"},{"tid":2353967,"ctgr":"오피스텔(주거)","saNo":"2024-63912","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 천호동 167-73, 6층 607호 (천호동,에스아이팰리스)","apslAmt":216000000,"minbAmt":110592000,"statNm":"유찰 3회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2373454,"ctgr":"도시형생활주택","saNo":"2025-154","crtDpt":"서울동부4계","regnAdrs":"서울 강동구 천호동 291-11, 3층 301호 (천호동,우노펠리스5차) 외 1필지","apslAmt":362000000,"minbAmt":231680000,"statNm":"유찰 2회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 3~4억"},{"tid":2368484,"ctgr":"도시형생활주택","saNo":"2025-50081","crtDpt":"서울동부4계","regnAdrs":"서울 광진구 중곡동 240-10, 3층301호 (중곡동,프라임하우스) 외 2필지","apslAmt":453000000,"minbAmt":362400000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 3~4억"},{"tid":2377533,"ctgr":"다세대주택","saNo":"2025-50988","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 능동 177-3 빌라헤르메스 지층103호 외 2필지","apslAmt":120000000,"minbAmt":76800000,"statNm":"유찰 2회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2377542,"ctgr":"도시형생활주택","saNo":"2025-51075","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 길동 369-2 새봄빌라 4층401호 외 1필지","apslAmt":230000000,"minbAmt":230000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2390270,"ctgr":"도시형생활주택","saNo":"2025-51129","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 구의동 80-29 구의동 홀가하우스 8층802호 외 2필지","apslAmt":364000000,"minbAmt":291200000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2491038,"ctgr":"도시형생활주택","saNo":"2025-51173","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 길동 457-1 강동렘브란트 11층1106호","apslAmt":126000000,"minbAmt":126000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1억이하"},{"tid":2390273,"ctgr":"다세대주택","saNo":"2025-51200","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 중곡동 239-54 미래하이츠 4층401호 외 1필지","apslAmt":635638000,"minbAmt":508510000,"statNm":"유찰 1회","bidDt":"26.02.02","splCdtn":"위반건축물,임차권등기,대항력 있는 임차인,공시가 3~4억"},{"tid":2395845,"ctgr":"다세대주택","saNo":"2025-51249","crtDpt":"서울동부4계","regnAdrs":"서울특별시 성동구 사근동 230-1 청계더하임 101동 5층504호","apslAmt":279000000,"minbAmt":279000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2397947,"ctgr":"다세대주택","saNo":"2025-51261","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 중곡동 247-21 한양팰리스 5층 504호 외 3필지","apslAmt":300000000,"minbAmt":300000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2395378,"ctgr":"다세대주택","saNo":"2025-51272","crtDpt":"서울동부4계","regnAdrs":"서울특별시 송파구 마천동 306-14 트라움바우 4층 401호","apslAmt":451000000,"minbAmt":451000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2396352,"ctgr":"다가구주택","saNo":"2025-51280","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 암사동 470-12 엘림파크","apslAmt":1996358400,"minbAmt":1996358400,"statNm":"신건","bidDt":"26.02.02","splCdtn":"위반건축물,임차권등기"},{"tid":2397948,"ctgr":"오피스텔(주거)","saNo":"2025-51301","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 길동 415-16 청광플러스원큐브3차 2층207호","apslAmt":111000000,"minbAmt":111000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2398829,"ctgr":"도시형생활주택","saNo":"2025-51323","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 천호동 169-9 다성이즈빌 9층901호 외 1필지","apslAmt":179000000,"minbAmt":179000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2402142,"ctgr":"단독주택","saNo":"2025-51373","crtDpt":"서울동부4계","regnAdrs":"서울특별시 광진구 군자동 60-28","apslAmt":866507600,"minbAmt":866507600,"statNm":"신건","bidDt":"26.02.02","splCdtn":"위반건축물,임차권등기,공시가 3~4억"},{"tid":2404401,"ctgr":"도시형생활주택","saNo":"2025-51401","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 길동 370-16 아성플러스 3층301호","apslAmt":309000000,"minbAmt":309000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2406247,"ctgr":"도시형생활주택","saNo":"2025-51454","crtDpt":"서울동부4계","regnAdrs":"서울특별시 강동구 천호동 180-15 아띠랑스에코타운 3층303호","apslAmt":332000000,"minbAmt":332000000,"statNm":"신건","bidDt":"26.02.02","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2147724,"ctgr":"다세대주택","saNo":"2022-56943","crtDpt":"서울서부7계","regnAdrs":"서울 은평구 응암동 578-33, 5층 503호 (응암동,더세영) 외 1필지","apslAmt":296000000,"minbAmt":296000000,"statNm":"유찰 12회","bidDt":"26.02.03","splCdtn":"선순위 가등기,임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2199131,"ctgr":"다세대주택","saNo":"2023-55183","crtDpt":"서울서부7계","regnAdrs":"서울 은평구 갈현동 472-6, 5층 502호","apslAmt":184000000,"minbAmt":19757000,"statNm":"유찰 10회","bidDt":"26.02.03","splCdtn":"임차인우선매수신고,임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2212462,"ctgr":"연립주택","saNo":"2023-56803","crtDpt":"서울서부5계","regnAdrs":"서울 서대문구 남가좌동 364-1, 1층101호 외 3필지","apslAmt":236000000,"minbAmt":20272000,"statNm":"유찰 11회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2220378,"ctgr":"다세대주택","saNo":"2023-57738","crtDpt":"서울서부5계","regnAdrs":"서울 은평구 구산동 177-132, 6층602호 (구산동,프라임하우스) 외 1필지","apslAmt":210000000,"minbAmt":4729000,"statNm":"유찰 17회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2222397,"ctgr":"다세대주택","saNo":"2023-58175","crtDpt":"서울서부5계","regnAdrs":"서울 은평구 신사동 10-14, 2층202호 (신사동,해창위너스빌)","apslAmt":318000000,"minbAmt":13986000,"statNm":"유찰 14회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2232316,"ctgr":"다세대주택","saNo":"2023-59529","crtDpt":"서울서부7계","regnAdrs":"서울 서대문구 홍은동 204-7, 5층402호 (홍은동,에코빌) 외 1필지","apslAmt":192000000,"minbAmt":25770000,"statNm":"유찰 9회","bidDt":"26.02.03","splCdtn":"위반건축물,임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2226786,"ctgr":"도시형생활주택","saNo":"2023-120293","crtDpt":"서울남부12계","regnAdrs":"서울 강서구 화곡동 504-51, 제6층 제603호 (화곡동, 서림그랑빌)","apslAmt":243000000,"minbAmt":79626000,"statNm":"유찰 5회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2278155,"ctgr":"오피스텔(주거)","saNo":"2024-944","crtDpt":"서울서부7계","regnAdrs":"서울 은평구 불광동 281-110, 비동 10층1002호 (불광동,샹그리라)","apslAmt":251000000,"minbAmt":128512000,"statNm":"유찰 3회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 2~3억"},{"tid":2281768,"ctgr":"도시형생활주택","saNo":"2024-1856","crtDpt":"서울남부12계","regnAdrs":"서울 구로구 구로동 44-5, 제6층 제611호 (구로동, 비즈트위트레인보우)","apslAmt":112000000,"minbAmt":89600000,"statNm":"유찰 1회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2338245,"ctgr":"단독주택","saNo":"2024-2599","crtDpt":"서울서부5계","regnAdrs":"서울특별시 용산구 한남동 568-109 외 3필지","apslAmt":3403611800,"minbAmt":3403611800,"statNm":"신건","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,NPL물건"},{"tid":2295437,"ctgr":"다세대주택","saNo":"2024-55432","crtDpt":"서울서부7계","regnAdrs":"서울 은평구 구산동 374, 에이동 3층302호 (구산동,야긴하이빌)","apslAmt":290000000,"minbAmt":24910000,"statNm":"유찰 11회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"},{"tid":2295019,"ctgr":"오피스텔(주거)","saNo":"2024-55586","crtDpt":"서울서부7계","regnAdrs":"서울 마포구 상암동 1734, 7층에이716호 (상암동,상암한화오벨리스크)","apslAmt":106000000,"minbAmt":11382000,"statNm":"유찰 10회","bidDt":"26.02.03","splCdtn":"임차권등기,대항력 있는 임차인,공시가 1~2억"}]}

# 세션 설정
session = requests.Session()
session.headers.update({
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8',
    'Accept-Language': 'ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7',
    'Connection': 'keep-alive',
    'Referer': 'https://tankauction.com/ca/caList.php?page=1',
    'Cookie': COOKIE_STRING,
})


def get_tenant_info(tid, total_count=100):
    """상세 페이지에서 임차인 현황 스크래핑"""
    url = f'https://tankauction.com/ca/caView.php?tid={tid}&chkNo=1&TotNo={total_count}'

    try:
        response = session.get(url, timeout=30)
        if response.status_code != 200:
            return None
    except Exception as e:
        print(f"  요청 오류: {e}")
        return None

    soup = BeautifulSoup(response.text, 'html.parser')

    tenant_section = soup.find('div', {'id': 'lyCnt_leas'})
    if not tenant_section:
        return {'말소기준일': '', '소액기준일': '', '배당요구종기일': '', '임차인목록': [], '기타사항': ''}

    result = {'말소기준일': '', '소액기준일': '', '배당요구종기일': '', '임차인목록': [], '기타사항': ''}

    span_box = tenant_section.find('span', class_='spanBox')
    if span_box:
        text = span_box.get_text()
        malso = re.search(r'말소기준일\s*:\s*(\d{4}-\d{2}-\d{2})', text)
        soaek = re.search(r'소액기준일\s*:\s*(\d{4}-\d{2}-\d{2})', text)
        badang = re.search(r'배당요구종기일\s*:\s*(\d{4}-\d{2}-\d{2})', text)
        if malso: result['말소기준일'] = malso.group(1)
        if soaek: result['소액기준일'] = soaek.group(1)
        if badang: result['배당요구종기일'] = badang.group(1)

    tenant_rows = tenant_section.find_all('tr', class_='dtLeasTr')
    for row in tenant_rows:
        tds = row.find_all('td')
        if len(tds) >= 8:
            tenant = {
                '점유목록': tds[0].get_text(strip=True),
                '임차인': tds[1].get_text(strip=True),
                '점유부분_기간': tds[2].get_text(separator=' ', strip=True),
                '전입_확정_배당': tds[3].get_text(separator=' | ', strip=True),
                '보증금_차임': tds[4].get_text(strip=True),
                '대항력': tds[5].get_text(strip=True),
                '분석': tds[6].get_text(separator=' ', strip=True),
                '기타': tds[7].get_text(strip=True)
            }
            result['임차인목록'].append(tenant)

    btbl = tenant_section.find('table', class_='Btbl_list')
    if btbl:
        etc_td = btbl.find('td')
        if etc_td:
            result['기타사항'] = etc_td.get_text(separator=' ', strip=True)

    return result


def create_excel(all_results, output_path):
    """엑셀 파일 생성"""
    df = pd.DataFrame(all_results)
    wb = Workbook()
    ws = wb.active
    ws.title = '임차인현황'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    headers = list(df.columns)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for row_idx, row_data in df.iterrows():
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx + 2, column=col_idx, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if headers[col_idx-1] == '대항력' and value == '있음':
                cell.font = Font(color='FF0000', bold=True)

    widths = {'A':10,'B':15,'C':12,'D':50,'E':15,'F':15,'G':10,'H':12,'I':30,'J':12,'K':12,'L':12,'M':8,'N':12,'O':25,'P':30,'Q':20,'R':10,'S':35,'T':40,'U':60}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = 'A2'
    wb.save(output_path)


def main():
    print("=" * 60)
    print("탱크옥션 임차인 현황 스크래핑")
    print("=" * 60)

    # 내장된 JSON 데이터 사용
    items = AUCTION_DATA['item']
    total_count = len(items)
    print(f"총 {total_count}개 물건 처리\n")

    all_results = []

    for idx, item in enumerate(items, 1):
        tid = item['tid']
        sa_no = item.get('saNo', '')
        addr = item.get('regnAdrs', '')
        category = item.get('ctgr', '')
        apsl_amt = item.get('apslAmt', 0)
        minb_amt = item.get('minbAmt', 0)
        stat_nm = item.get('statNm', '')
        bid_dt = item.get('bidDt', '')
        spl_cdtn = item.get('splCdtn', '')

        print(f"[{idx}/{total_count}] {sa_no} (tid={tid})", end=" ")

        tenant_info = get_tenant_info(tid, total_count)

        if tenant_info and tenant_info.get('임차인목록'):
            for tenant in tenant_info['임차인목록']:
                row = {
                    'TID': tid, '사건번호': sa_no, '물건종류': category, '소재지': addr,
                    '감정가': apsl_amt, '최저가': minb_amt, '진행상태': stat_nm,
                    '입찰일': bid_dt, '매각조건': spl_cdtn,
                    '말소기준일': tenant_info.get('말소기준일', ''),
                    '소액기준일': tenant_info.get('소액기준일', ''),
                    '배당요구종기일': tenant_info.get('배당요구종기일', ''),
                    '점유목록': tenant.get('점유목록', ''),
                    '임차인': tenant.get('임차인', ''),
                    '점유부분_기간': tenant.get('점유부분_기간', ''),
                    '전입_확정_배당': tenant.get('전입_확정_배당', ''),
                    '보증금_차임': tenant.get('보증금_차임', ''),
                    '대항력': tenant.get('대항력', ''),
                    '분석': tenant.get('분석', ''),
                    '임차인기타': tenant.get('기타', ''),
                    '기타사항': tenant_info.get('기타사항', '')
                }
                all_results.append(row)
            print(f"-> {len(tenant_info['임차인목록'])}명")
        else:
            row = {
                'TID': tid, '사건번호': sa_no, '물건종류': category, '소재지': addr,
                '감정가': apsl_amt, '최저가': minb_amt, '진행상태': stat_nm,
                '입찰일': bid_dt, '매각조건': spl_cdtn,
                '말소기준일': tenant_info.get('말소기준일', '') if tenant_info else '',
                '소액기준일': tenant_info.get('소액기준일', '') if tenant_info else '',
                '배당요구종기일': tenant_info.get('배당요구종기일', '') if tenant_info else '',
                '점유목록': '', '임차인': '정보없음', '점유부분_기간': '',
                '전입_확정_배당': '', '보증금_차임': '', '대항력': '',
                '분석': '', '임차인기타': '', '기타사항': ''
            }
            all_results.append(row)
            print("-> 정보없음")

        time.sleep(0.3)

    output_path = '탱크옥션_임차인현황.xlsx'
    create_excel(all_results, output_path)

    print("\n" + "=" * 60)
    print(f"완료! 총 {len(all_results)}개 레코드 저장")
    print(f"파일: {output_path}")
    print("=" * 60)


if __name__ == '__main__':
    main()